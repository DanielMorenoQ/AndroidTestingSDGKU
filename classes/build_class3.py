"""Generate the Class 3 slide deck: Stubs, Mocks & Dependency Injection."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- Theme -------------------------------------------------------------------
NAVY = RGBColor(0x0D, 0x1B, 0x2A)
BLUE = RGBColor(0x1B, 0x4D, 0x7A)
ACCENT = RGBColor(0x3D, 0xDC, 0x84)
RED = RGBColor(0xE5, 0x3E, 0x3E)
AMBER = RGBColor(0xF4, 0xB4, 0x00)
PURPLE = RGBColor(0x6C, 0x5C, 0xE7)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
GREY = RGBColor(0x5A, 0x6A, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG = RGBColor(0x1E, 0x1E, 0x2E)
CODE_FG = RGBColor(0xE6, 0xE6, 0xE6)

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def box(slide, l, t, w, h):
    return slide.shapes.add_textbox(l, t, w, h).text_frame


def set_run(run, text, size, color, bold=False, italic=False, font=BODY_FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def accent_bar(slide):
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.18), SH)
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()


def title_bar(slide, number, title):
    accent_bar(slide)
    tf = box(slide, Inches(0.6), Inches(0.35), Inches(11.8), Inches(1.0))
    set_run(tf.paragraphs[0].add_run(), title, 32, NAVY, bold=True)
    chip = slide.shapes.add_shape(1, Inches(12.3), Inches(0.4), Inches(0.7), Inches(0.55))
    chip.fill.solid(); chip.fill.fore_color.rgb = BLUE; chip.line.fill.background()
    chip.text_frame.word_wrap = False
    cp = chip.text_frame.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
    set_run(cp.add_run(), str(number), 18, WHITE, bold=True)
    ln = slide.shapes.add_shape(1, Inches(0.6), Inches(1.3), Inches(2.4), Inches(0.06))
    ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()


def bullets(slide, items, top=Inches(1.7), left=Inches(0.8), width=Inches(11.7),
            height=Inches(5.2), size=20):
    tf = box(slide, left, top, width, height)
    tf.word_wrap = True
    first = True
    for item in items:
        text, level = item if isinstance(item, tuple) else (item, 0)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(10)
        bullet = "•  " if level == 0 else "–  "
        set_run(p.add_run(), bullet + text, size - level * 2, NAVY if level == 0 else GREY)
    return tf


def speaker(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def shp(slide, kind, l, t, w, h, fill, line=None, line_w=1.0):
    sp = slide.shapes.add_shape(kind, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is not None:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def label(sp, text, size, color, bold=True, sub=None, sub_color=None):
    tf = sp.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), text, size, color, bold=bold)
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        set_run(p2.add_run(), sub, size - 8, sub_color or color)


def code_slide(number, title, code, size=14):
    slide = add_slide(); bg(slide, NAVY); accent_bar(slide)
    tf = box(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.9))
    set_run(tf.paragraphs[0].add_run(), title, 28, WHITE, bold=True)
    chip = slide.shapes.add_shape(1, Inches(12.3), Inches(0.35), Inches(0.7), Inches(0.55))
    chip.fill.solid(); chip.fill.fore_color.rgb = ACCENT; chip.line.fill.background()
    chip.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_run(chip.text_frame.paragraphs[0].add_run(), str(number), 18, NAVY, bold=True)
    panel = slide.shapes.add_shape(1, Inches(0.6), Inches(1.35), Inches(12.1), Inches(5.7))
    panel.fill.solid(); panel.fill.fore_color.rgb = CODE_BG
    panel.line.color.rgb = BLUE; panel.line.width = Pt(1)
    ctf = panel.text_frame; ctf.word_wrap = True; ctf.vertical_anchor = MSO_ANCHOR.TOP
    ctf.margin_left = Inches(0.3); ctf.margin_top = Inches(0.2)
    ctf.margin_right = Inches(0.2); ctf.margin_bottom = Inches(0.2)
    for i, line in enumerate(code.split("\n")):
        p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
        p.line_spacing = 1.0
        color = RGBColor(0x7E, 0xC6, 0x99) if line.strip().startswith("//") else CODE_FG
        set_run(p.add_run(), line if line else " ", size, color, font=CODE_FONT)
    return slide


def table_slide(number, title, headers, rows, col_widths=None, note=None, fsize=13):
    slide = add_slide(); bg(slide, WHITE); title_bar(slide, number, title)
    nrows = len(rows) + 1
    gfx = slide.shapes.add_table(nrows, len(headers), Inches(0.7), Inches(1.7),
                                 Inches(11.9), Inches(0.6) * nrows)
    table = gfx.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
        set_run(cell.text_frame.paragraphs[0].add_run(), h, 15, WHITE, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
            set_run(cell.text_frame.paragraphs[0].add_run(), val, fsize, NAVY)
    if note:
        tf = box(slide, Inches(0.7), Inches(6.7), Inches(11.9), Inches(0.6))
        set_run(tf.paragraphs[0].add_run(), note, 14, GREY, italic=True)
    return slide


def doubles_spectrum(slide, top=Inches(5.2)):
    """Dummy -> Stub -> Fake -> Spy -> Mock, simple to behavioural."""
    steps = [("Dummy", GREY), ("Stub", BLUE), ("Fake", ACCENT),
             ("Spy", AMBER), ("Mock", PURPLE)]
    x = Inches(0.9); w = Inches(2.15); gap = Inches(0.25); h = Inches(1.0)
    for t, c in steps:
        b = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, top, w, h, c)
        label(b, t, 18, NAVY if c in (ACCENT, AMBER) else WHITE)
        x = Emu(int(x) + int(w) + int(gap))
    cap = box(slide, Inches(0.9), Emu(int(top) + int(h) + int(Inches(0.1))), Inches(11.6), Inches(0.5))
    cap.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_run(cap.paragraphs[0].add_run(),
            "just data  ←——————————————————————→  verifies behaviour", 14, GREY, italic=True)


def di_before_after(slide, top=Inches(2.0)):
    # BEFORE: tight coupling
    cap1 = box(slide, Inches(0.9), Emu(int(top) - int(Inches(0.5))), Inches(5.3), Inches(0.5))
    set_run(cap1.paragraphs[0].add_run(), "Before — welded to a concrete class", 16, RED, bold=True)
    svc = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), top, Inches(3.4), Inches(1.0), BLUE)
    label(svc, "LoginService", 15, WHITE)
    store = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), Emu(int(top) + int(Inches(2.0))),
                Inches(3.4), Inches(1.0), NAVY)
    label(store, "new UserStore()", 14, WHITE)
    shp(slide, MSO_SHAPE.DOWN_ARROW, Inches(3.0), Emu(int(top) + int(Inches(1.1))),
        Inches(0.6), Inches(0.8), RED)
    # AFTER: depends on interface
    cap2 = box(slide, Inches(7.3), Emu(int(top) - int(Inches(0.5))), Inches(5.3), Inches(0.5))
    set_run(cap2.paragraphs[0].add_run(), "After — depends on an interface", 16, ACCENT, bold=True)
    svc2 = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), top, Inches(3.6), Inches(1.0), BLUE)
    label(svc2, "LoginService", 15, WHITE)
    iface = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Emu(int(top) + int(Inches(1.7))),
                Inches(3.6), Inches(0.9), ACCENT)
    label(iface, "«interface»  UserRepository", 13, NAVY)
    shp(slide, MSO_SHAPE.DOWN_ARROW, Inches(9.5), Emu(int(top) + int(Inches(1.05))),
        Inches(0.6), Inches(0.6), GREY)
    real = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Emu(int(top) + int(Inches(3.1))),
               Inches(2.3), Inches(0.85), NAVY)
    label(real, "UserStore", 13, WHITE)
    fake = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.0), Emu(int(top) + int(Inches(3.1))),
               Inches(2.3), Inches(0.85), PURPLE)
    label(fake, "Fake / Mock", 13, WHITE)
    for cx in (Inches(8.35), Inches(11.15)):
        conn = slide.shapes.add_connector(2, Inches(9.8), Emu(int(top) + int(Inches(2.6))),
                                          cx, Emu(int(top) + int(Inches(3.1))))
        conn.line.color.rgb = GREY; conn.line.width = Pt(1.25)


# ============================================================================
# Slide 1 — Title
# ============================================================================
s = add_slide(); bg(s, NAVY)
blk = s.shapes.add_shape(1, 0, Inches(2.7), SW, Inches(0.12))
blk.fill.solid(); blk.fill.fore_color.rgb = ACCENT; blk.line.fill.background()
tf = box(s, Inches(1.0), Inches(1.5), Inches(11.3), Inches(1.4))
set_run(tf.paragraphs[0].add_run(), "Stubs, Mocks & Dependency Injection", 44, WHITE, bold=True)
tf2 = box(s, Inches(1.0), Inches(2.9), Inches(11.3), Inches(0.9))
set_run(tf2.paragraphs[0].add_run(), "Class 3 — Test doubles and how DI enables them", 25, ACCENT, bold=True)
tf3 = box(s, Inches(1.0), Inches(4.0), Inches(11.3), Inches(1.2))
set_run(tf3.paragraphs[0].add_run(),
        "Dummy · Stub · Fake · Spy · Mock  ·  MockK  ·  constructor injection", 20, LIGHT)
speaker(s, "Class 3 is about isolating the class under test. We learn the five test doubles, "
           "when to use each, and the Dependency Injection principle that makes them possible.")

# ============================================================================
# Slide 2 — Where we are
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 2, "Where we are")
bullets(s, [
    ("1.  Unit, integration and UI tests", 0),
    ("2.  TDD", 0),
    ("3.  Mocks, stubs and fakes  ◀  today", 0),
    ("4.  UI testing and CI/CD", 0),
    ("5.  Fixing flaky tests", 0),
    ("6.  Catching bugs, plus the exam", 0),
])
speaker(s, "We can test logic and UI. Today: how to isolate a class from its real collaborators "
           "so tests stay fast and deterministic.")

# ============================================================================
# Slide 3 — Agenda
# ============================================================================
table_slide(3, "Today — 3 hours",
            ["Time", "Block", "Activity"],
            [
                ["0:10–0:50", "Theory", "Test doubles + DI"],
                ["0:50–1:05", "Live demo", "Refactor LoginService behind an interface + a fake"],
                ["1:15–2:00", "Exercise 1", "Fakes, stubs & MockK mocks on the auth layer"],
                ["2:00–2:45", "Exercise 2", "ShopViewModel + ProductRepository with doubles"],
                ["2:45–3:00", "Wrap-up", "Compare doubles, pitfalls, homework"],
            ],
            col_widths=[Inches(2.4), Inches(2.3), Inches(7.2)])

# ============================================================================
# Slide 4 — Why test doubles
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 4, "Why fake a collaborator?")
bullets(s, [
    "Real collaborators can be slow (network, database, disk)",
    "Non-deterministic: time, randomness, the network being down",
    "Have side effects you don't want in a test (emails, payments)",
    "Hard to force into an error state on demand",
    "A double gives you a fast, predictable, controllable stand-in",
])

# ============================================================================
# Slide 5 — What is a test double
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 5, "What is a test double?")
bullets(s, [
    "A stand-in for a real collaborator of the class under test",
    "SUT = System Under Test (the class you're testing)",
    "Collaborator = a dependency the SUT talks to",
    "The double replaces the collaborator so the test controls it",
    "Term coined by Gerard Meszaros; taxonomy popularised by Martin Fowler",
])

# ============================================================================
# Slide 6 — The five doubles (spectrum)
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 6, "The five doubles")
bullets(s, [
    "Dummy — passed to fill a parameter, never actually used",
    "Stub — returns canned answers to calls",
    "Fake — a real but lightweight implementation (in-memory)",
    "Spy — a stub that also records how it was called",
    "Mock — pre-programmed with expectations, verifies behaviour",
], height=Inches(3.0))
doubles_spectrum(s)

# ============================================================================
# Slides 7-11 — each double
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 7, "Dummy")
bullets(s, [
    "An object passed only to satisfy a parameter list",
    "Never used by the code path under test",
    "Often just null, a no-op, or an empty object",
    "Example: a Logger you must pass but the test never inspects",
])

s = add_slide(); bg(s, WHITE); title_bar(s, 8, "Stub")
bullets(s, [
    "Provides canned answers to the calls made during the test",
    "No logic, no verification — just data in",
    "You assert on the SUT's result → state verification",
    'MockK: every { repo.passwordFor(any()) } returns "password123"',
])

s = add_slide(); bg(s, WHITE); title_bar(s, 9, "Fake")
bullets(s, [
    "A working implementation — but not production-ready",
    "Classic example: an in-memory repository instead of a real database",
    "Reads naturally and is reusable across many tests",
    "Often the clearest double: it's just a small class",
])

s = add_slide(); bg(s, WHITE); title_bar(s, 10, "Spy")
bullets(s, [
    "A stub (or real object) that also records how it was called",
    "Lets you keep real behaviour and still verify interactions",
    "Useful to wrap an existing fake and assert on its calls",
    "MockK: spyk(FakeUserRepository())",
])

s = add_slide(); bg(s, WHITE); title_bar(s, 11, "Mock")
bullets(s, [
    "Pre-programmed with expectations about the calls it should receive",
    "The test fails if those interactions don't happen as specified",
    "You assert the SUT talked to it correctly → behaviour verification",
    "MockK: verify(exactly = 1) { repo.passwordFor(\"ana@example.com\") }",
])

# ============================================================================
# Slide 12 — Stub vs Mock
# ============================================================================
table_slide(12, "Stub vs Mock — the key distinction",
            ["", "Stub", "Mock"],
            [
                ["Direction", "Data flows IN to the SUT", "Calls flow OUT of the SUT"],
                ["You assert", "SUT state / return value", "That the call happened"],
                ["Style", "State verification", "Interaction verification"],
                ["Breaks when", "The result is wrong", "The collaboration changes"],
                ["Prefer when", "You need data", "The interaction IS the behaviour"],
            ],
            col_widths=[Inches(2.6), Inches(4.65), Inches(4.65)])

# ============================================================================
# Slide 13 — MockK cheat sheet (code)
# ============================================================================
code_slide(13, "MockK cheat sheet", '''val repo = mockk<UserRepository>()                  // strict by default

// STUB a return value
every { repo.passwordFor("a@b.com") } returns "secret"
every { repo.passwordFor(any()) } throws IOException()

// VERIFY an interaction (this is what makes it a mock)
verify { repo.passwordFor("a@b.com") }
verify(exactly = 0) { repo.passwordFor(any()) }     // was NOT called
confirmVerified(repo)                               // nothing else happened

val relaxed = mockk<UserRepository>(relaxed = true) // auto-returns defaults
val spy = spyk(FakeUserRepository())                // real behaviour + recording

// suspend functions
coEvery { api.passwordFor(any()) } returns "secret"
coVerify { api.passwordFor(any()) }''')

# ============================================================================
# Slide 14 — Choosing a double
# ============================================================================
table_slide(14, "Which double should I reach for?",
            ["You need to…", "Use"],
            [
                ["Fill a parameter you never use", "Dummy"],
                ["Feed canned data and check a result", "Stub"],
                ["Replace a DB/network with something that really works", "Fake"],
                ["Keep real behaviour but verify a call", "Spy"],
                ["Prove the SUT called a collaborator correctly", "Mock"],
            ],
            col_widths=[Inches(8.4), Inches(3.5)],
            note="Rule of thumb: prefer fakes/stubs; reach for a mock when the interaction is the behaviour.")

# ============================================================================
# Slide 15 — DI principle
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 15, "Dependency Injection")
bullets(s, [
    "A class receives its dependencies from outside instead of creating them",
    "Most common form: pass them through the constructor",
    "Production passes the real thing; tests pass a double",
    "No framework required — DI is a principle, not a library",
])

# ============================================================================
# Slide 16 — DIP / IoC
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 16, "Depend on abstractions")
bullets(s, [
    "Dependency Inversion Principle: depend on an interface, not a concrete class",
    "High-level code (LoginService) shouldn't know the low-level detail (UserStore)",
    "Both depend on the abstraction (UserRepository)",
    "Inversion of Control: who creates the dependency moves up and out",
    "The result: collaborators become swappable — including for doubles",
])

# ============================================================================
# Slide 17 — Constructor injection (diagram)
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 17, "Constructor injection")
di_before_after(s)

# ============================================================================
# Slide 18 — Ways to inject
# ============================================================================
table_slide(18, "Ways to inject dependencies",
            ["Approach", "What it is", "When"],
            [
                ["Manual / constructor", "Pass dependencies by hand", "Small apps, this course"],
                ["Composition root", "One place wires the graph (the Activity)", "Keeps wiring out of logic"],
                ["Service Locator", "A registry you ask for dependencies", "Simple, but hides deps"],
                ["Hilt / Dagger", "Annotation-based DI framework", "Larger Android apps"],
            ],
            col_widths=[Inches(3.1), Inches(5.6), Inches(3.2)])

# ============================================================================
# Slide 19 — DI makes doubles trivial (code)
# ============================================================================
code_slide(19, "DI makes doubles trivial", '''interface UserRepository {
    fun passwordFor(email: String): String?
}

// Production implementation
class UserStore : UserRepository {
    private val users = mapOf("ana@example.com" to "password123")
    override fun passwordFor(email: String) = users[email]
}

// SUT depends on the ABSTRACTION, with a real default
class LoginService(private val users: UserRepository = UserStore()) {
    fun login(email: String, password: String): LoginResult { /* ... */ }
}

// A test hands it a fake — no network, no real store
val service = LoginService(FakeUserRepository().withUser("ana@example.com", "password123"))''')

# ============================================================================
# Slide 20 — Pitfalls
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 20, "Pitfalls & anti-patterns")
bullets(s, [
    "Over-mocking: tests mirror the implementation and break on every refactor",
    "Don't mock types you don't own — wrap the SDK behind your own interface first",
    "Don't mock value objects (Product, LoginResult) — just construct them",
    "A mock with no verification is really just a stub — name it honestly",
    "Too many mocks in one test is a design smell: the class has too many collaborators",
])
speaker(s, "The most common mistake is mocking everything. Prefer fakes and stubs; mock only at "
           "true boundaries like network, database and time.")

# ============================================================================
# Slide 21 — Summary & exercise
# ============================================================================
s = add_slide(); bg(s, NAVY); accent_bar(s)
tf = box(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.0))
set_run(tf.paragraphs[0].add_run(), "Summary & exercise", 32, WHITE, bold=True)
ln = s.shapes.add_shape(1, Inches(0.65), Inches(1.35), Inches(2.4), Inches(0.06))
ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()
tf = box(s, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0)); tf.word_wrap = True
items = [
    ("A test double stands in for a real collaborator", False),
    ("Stub = state verification; Mock = interaction verification", False),
    ("DI (constructor + interface) is what makes doubles possible", False),
    ("The composition root is the only place that knows concrete types", False),
    ("Exercise 1: fakes, stubs & MockK mocks on LoginService / UserRepository", True),
    ("Exercise 2: ShopViewModel + ProductRepository tested with doubles", True),
]
for i, (t, hw) in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(12)
    set_run(p.add_run(), "•  " + t, 21, ACCENT if hw else LIGHT, bold=hw)

prs.save("slides/Class3_Stubs_Mocks_DI.pptx")
print("Saved slides/Class3_Stubs_Mocks_DI.pptx with", len(prs.slides._sldIdLst), "slides")
