"""Generate Class 4 slides with instructor notes; run from any directory."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "slides" / "Class4_UI_Testing_CI_CD.pptx"
NAVY = RGBColor(0x0D, 0x1B, 0x2A)
BLUE = RGBColor(0x1B, 0x4D, 0x7A)
GREEN = RGBColor(0x3D, 0xDC, 0x84)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
GREY = RGBColor(0x5A, 0x6A, 0x7A)
CODE_BG = RGBColor(0x1E, 0x1E, 0x2E)


def textbox(slide, text, left, top, width, height, size=22,
            color=NAVY, bold=False, font="Calibri"):
    frame = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    ).text_frame
    frame.word_wrap = True
    for index, line in enumerate(text.splitlines()):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(12)
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return frame


def shape(slide, kind, left, top, width, height, color):
    result = slide.shapes.add_shape(
        kind, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    result.fill.solid()
    result.fill.fore_color.rgb = color
    result.line.fill.background()
    return result


def slide_base(deck, title, notes, dark=False):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY if dark else WHITE
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, 7.5, GREEN)
    textbox(slide, title, 0.6, 0.35, 11.4, 0.85, 30,
            WHITE if dark else NAVY, True)
    shape(slide, MSO_SHAPE.RECTANGLE, 0.6, 1.3, 2.4, 0.06, GREEN)
    frame = textbox(slide, str(len(deck.slides)), 12.2, 0.4, 0.7, 0.55,
                    18, GREEN if dark else BLUE, True)
    frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    textbox(slide, "Class 4 | UI Testing and CI/CD", 0.65, 7.15, 11, 0.25,
            11, LIGHT if dark else GREY)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def bullet_slide(deck, title, lines, notes, dark=False):
    slide = slide_base(deck, title, notes, dark)
    textbox(slide, "\n".join("- " + line for line in lines),
            0.8, 1.75, 11.7, 5.1, 23, LIGHT if dark else NAVY)


def code_slide(deck, title, code, notes):
    slide = slide_base(deck, title, notes, True)
    shape(slide, MSO_SHAPE.RECTANGLE, 0.65, 1.65, 12, 5.25, CODE_BG)
    frame = textbox(slide, code, 0.85, 1.82, 11.5, 4.9,
                    19, LIGHT, font="Consolas")
    for paragraph in frame.paragraphs:
        paragraph.space_after = Pt(2)


def flow_slide(deck, title, stages, notes):
    slide = slide_base(deck, title, notes)
    width = 2.55
    for index, (heading, detail) in enumerate(stages):
        left = 0.7 + index * 3.1
        shape(slide, MSO_SHAPE.RECTANGLE, left, 2.2, width, 1.1, BLUE)
        frame = textbox(slide, heading, left + 0.12, 2.4, width - 0.24,
                        0.75, 24, WHITE, True)
        frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        textbox(slide, detail, left, 3.6, width, 2.5, 20, GREY)
        if index < len(stages) - 1:
            shape(slide, MSO_SHAPE.RIGHT_ARROW, left + width + 0.07,
                  2.6, 0.4, 0.35, GREEN)


def build_deck():
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)

    bullet_slide(deck, "UI Testing and CI/CD", [
        "Class 4 | Three-hour workshop",
        "Build a Shop and Cart, then test real user journeys",
        "Run the test suite automatically on GitHub",
        "Instructor script: exercises/Class-4-UI-Testing-CI-CD.md",
    ], "0:00. Open the project and script. Today follows classes 1-3. "
       "The current ShopActivity has only a welcome message; build the UI before testing it.", True)

    bullet_slide(deck, "What learners will deliver", [
        "A catalog and Cart screen with navigation and shared state",
        "Tests for empty cart, repeated adds, removal and Back",
        "A real login-to-Shop test across XML and Compose",
        "A Git repository, pull request and two green CI checks",
        "Evidence that a regression blocks the pull request",
    ], "0:03. Read the acceptance criteria. Do not promise production checkout or a release pipeline.")

    bullet_slide(deck, "Today | Exactly 180 minutes", [
        "00-10: recap and environment preflight",
        "10-30: UI test design, semantics and navigation",
        "30-65: build Shop and Cart screens",
        "65-75: break",
        "75-115: UI test lab; 115-130: Git and pull request",
        "130-165: GitHub Actions; 165-180: failure drill and wrap-up",
    ], "0:05. Start CI early in its block. First emulator runs can be slow; use a preflighted instructor repo as fallback.")

    bullet_slide(deck, "Preflight before live coding", [
        "Run existing JVM tests and assemble the debug APK",
        "Confirm an API 35+ emulator boots and ADB sees it",
        "Gradle daemon: JDK 25; local test launcher: JDK 21",
        "This project compiles and targets SDK 37",
        "Compose navigation and UI-test dependencies already exist",
        "GitHub account, Git CLI and a repository ready for CI",
    ], "0:07. Use script commands. API 35 is a test runtime, not a replacement for compileSdk 37. "
       "Do not silently downgrade AGP or the wrapper if dependency resolution fails.")

    flow_slide(deck, "One user journey, two UI technologies", [
        ("Login", "XML views\nEspresso\nMainActivity"),
        ("Shop", "Compose catalog\nAdd products\nShopActivity"),
        ("Cart", "Compose route\nQuantity and remove\nShared cart state"),
        ("Back", "Catalog route\nCart survives navigation\nNo duplicate route"),
    ], "0:10. Login opens ShopActivity with an Intent. Shop-to-Cart navigation uses Compose NavHost. "
       "Distinguish activity navigation from in-activity routes.")

    bullet_slide(deck, "Test at the cheapest useful layer", [
        "JVM unit tests: pricing and validation rules",
        "Robolectric: existing host-side Android/Compose checks",
        "Compose UI tests: semantics, interactions and navigation",
        "Espresso: XML login fields and real activity transition",
        "A few full journeys complement many focused tests",
    ], "0:13. testDebugUnitTest includes existing Robolectric tests; connectedDebugAndroidTest needs a device. "
       "Do not claim that host-side tests prove emulator behavior.")

    bullet_slide(deck, "Acceptance criteria become assertions", [
        "New Shop: Cart (0), empty Cart message",
        "Add Keyboard twice: quantity 2, Cart (2), total $160.00",
        "Add Mouse too: two product rows, total $205.00",
        "Remove Keyboard: all of that row removed, total $45.00",
        "Cart -> Back -> Cart: remaining item is still present",
        "Invalid login stays put; valid login reaches catalog",
    ], "0:16. Totals are raw catalog totals, deliberately not ShoppingCartCalculator's discount rules. "
       "Remove means remove the whole product row, not decrement by one.")

    code_slide(deck, "Stable semantics, not screen coordinates", '''Modifier.testTag("product_list")
Modifier.testTag("add_keyboard")
Modifier.testTag("cart_count")
Modifier.testTag("quantity_keyboard")
Modifier.testTag("cart_total")

rule.onNodeWithTag("add_keyboard").performClick()
rule.onNodeWithTag("cart_count", useUnmergedTree = true)
    .assertTextEquals("Cart (1)")''', "0:19. Tags identify unique nodes; assertions inspect user-visible text. "
        "The count Text is inside a merging Button; query its tag in the unmerged tree. Click the Button's tag in the merged tree.")

    bullet_slide(deck, "Synchronization without sleeps", [
        "Compose actions/assertions synchronize with Compose work",
        "Use performScrollToNode for offscreen LazyColumn items",
        "Use idling resources for unmanaged Espresso async work",
        "For external async state, waitUntil with a bounded predicate",
        "Never replace a missing assertion with Thread.sleep",
    ], "0:23. Our in-memory synchronous fixture needs no network, sleeps or retries. "
       "A wait must name a condition; a longer timeout is not a diagnosis.")

    bullet_slide(deck, "Test fixture and ownership", [
        "Use a fixed local catalog; no accounts or real payments",
        "Store quantities above NavHost with rememberSaveable",
        "Navigate between shop and cart; both see the same state",
        "Launch a fresh Activity for each test",
        "Keep pricing scope explicit; reuse class-3 fakes for async bonus",
    ], "0:27. The fixture uses Keyboard $80, Mouse $45, Monitor $180, Headset $120 and Webcam $60. "
       "Saved state handles these small local values, not a durable cart backend.")

    bullet_slide(deck, "Build first | Catalog and Cart", [
        "Replace the placeholder ShopScreen with ShopApp",
        "Catalog: product rows, price and an Add command",
        "Cart: empty state, quantities, Remove and total",
        "Toolbar: Cart count on catalog, Back on Cart",
        "Wrap content in MaterialTheme; honor Scaffold insets",
        "Run manually before writing the UI tests",
    ], "0:30-0:40. Use the complete code in script section 3. Replace the old top-level placeholder, "
       "do not create duplicate ShopApp or ShopScreen definitions. Verify the five rows and buttons.")

    code_slide(deck, "NavHost owns routes, not cart lifetime", '''val nav = rememberNavController()
var quantities by rememberSaveable {
    mutableStateOf(hashMapOf<String, Int>())
}

NavHost(nav, startDestination = "shop") {
    composable("shop") { /* catalog */ }
    composable("cart") { /* cart */ }
}
// Toolbar action:
nav.navigate("cart") { launchSingleTop = true }
// Back action: nav.popBackStack()''', "0:40. Slide code is an excerpt; copy complete script code. "
        "Assign a NEW map after changes so Compose observes state, rather than mutating the existing HashMap.")

    bullet_slide(deck, "Checkpoint | Manual navigation lab", [
        "Log in; add Keyboard twice and Mouse once",
        "Open Cart: quantities 2 and 1; total $205.00",
        "Return to Shop, then reopen Cart: state unchanged",
        "Remove Keyboard: Mouse remains and total is $45.00",
        "Remove Mouse: empty Cart message, total $0.00",
        "Commit the built screens before test lab",
    ], "0:50-1:05. Give pairs time to implement and inspect small-screen scrolling. "
       "Ask where state lives. Take the break at 1:05 even if someone needs the reference code.")

    bullet_slide(deck, "Break | Resume with testable screens", [
        "10 minutes: 1:05-1:15",
        "Checkpoint: catalog, Cart navigation and stable tags work",
        "Next: tests that prove behavior, not just screen existence",
    ], "1:05. Leave this slide visible. Instructor opens the test source set and confirms emulator is available.", True)

    code_slide(deck, "Activity-hosted Compose UI test", '''@RunWith(AndroidJUnit4::class)
class ShopNavigationTest {
    @get:Rule
    val rule = createAndroidComposeRule<ShopActivity>()

    @Test fun newCart_isEmpty() {
        rule.onNodeWithTag("open_cart").performClick()
        rule.onNodeWithTag("empty_cart")
            .assertIsDisplayed()
        rule.onNodeWithTag("cart_total")
            .assertTextEquals("Total: $0.00")
    }
}''', "1:15. Save in androidTest, not test. Complete imports and five tests are in the script. "
        "The rule launches real ShopActivity and its navigation; do not call setContent again.")

    code_slide(deck, "Repeated add: prove quantity and price", '''rule.onNodeWithTag("product_list")
    .performScrollToNode(hasTestTag("add_keyboard"))
rule.onNodeWithTag("add_keyboard").performClick()
rule.onNodeWithTag("add_keyboard").performClick()
rule.onNodeWithTag("cart_count", useUnmergedTree = true)
    .assertTextEquals("Cart (2)")
rule.onNodeWithTag("open_cart").performClick()
rule.onNodeWithTag("quantity_keyboard")
    .assertTextEquals("Quantity: 2")
rule.onNodeWithTag("cart_total")
    .assertTextEquals("Total: $160.00")''', "1:22. Count alone could pass while the cart data is broken. "
        "Inspect both row quantity and total. Each test starts fresh; no ordering assumptions.")

    bullet_slide(deck, "Advanced lab | Navigation and state", [
        "Test 1: empty cart and zero total",
        "Test 2: repeated product increments quantity, not rows",
        "Test 3: two products survive Back and reopening Cart",
        "Test 4: remove all quantities; assert row absence and zero",
        "Test 5: scroll to Webcam, add, then verify Cart row",
        "Challenge: system Back, rotation and restored state",
    ], "1:27-1:40. Use pair work: one writes action, one names a discriminating assertion. "
       "System Back and recreation are extension tasks, not claims made by the baseline toolbar-Back test.")

    code_slide(deck, "XML login -> real Compose Shop", '''@get:Rule
val rule = createAndroidComposeRule<MainActivity>()

onView(withId(R.id.emailInput))
    .perform(replaceText("ana@example.com"))
onView(withId(R.id.passwordInput))
    .perform(replaceText("password123"), closeSoftKeyboard())
onView(withId(R.id.loginButton)).perform(click())

rule.onNodeWithTag("shop_title").assertIsDisplayed()
rule.onNodeWithTag("add_keyboard").performClick()
rule.onNodeWithTag("cart_count", useUnmergedTree = true)
    .assertTextEquals("Cart (1)")''', "1:40. Compose's test environment observes the launched ShopActivity root. "
        "Do not stub the Shop intent in this journey. Existing intent verification tests can remain as narrower checks.")

    bullet_slide(deck, "Run locally and inspect a real failure", [
        "./gradlew :app:testDebugUnitTest",
        "./gradlew :app:connectedDebugAndroidTest",
        "Single class: use instrumentation runner class argument",
        "First failure: read expected/actual and locate the screen",
        "Save reports; distinguish assertion, build and emulator failure",
    ], "1:48-1:55. Show script commands, instrumented HTML report and JVM XML. "
       "A provisioning error is not a failed product assertion. Do not rewrite unrelated existing tests.")

    bullet_slide(deck, "Git lab | Commit and open a pull request", [
        "This folder currently has no Git metadata",
        "Initialize main, inspect ignores, commit the baseline",
        "Never commit local.properties, build outputs or credentials",
        "Create class4/ui-ci branch; commit screens, tests, workflow",
        "Add GitHub remote, push branch, open a pull request",
        "Use exact named CI checks as merge requirements",
    ], "1:55-2:10. A remote URL alone does not authenticate Git. "
       "Use GitHub CLI browser login or the user's existing SSH/credential setup, never put tokens in the script.")

    flow_slide(deck, "CI | Two required checks", [
        ("Push / PR", "Checkout wrapper\nInstall JDK 25 + 21\nInstall Android SDK"),
        ("JVM + lint", "testDebugUnitTest\nlintDebug\nassembleDebug"),
        ("Device UI", "KVM emulator\nAPI 35 x86_64\nconnected tests"),
        ("Evidence", "Reports on failure\nDebug APK\nReview and merge"),
    ], "2:10. Jobs run independently, so fast JVM feedback does not wait for the emulator. "
       "Diagram order explains concepts; the YAML runs the two jobs in parallel, not sequentially.")

    bullet_slide(deck, "Why a runner needs two JDKs", [
        "Committed daemon criteria explicitly request Java 25",
        "Existing Test tasks explicitly request a Java 21 launcher",
        "Install 21 first; install 25 last so JAVA_HOME selects 25",
        "Expose both through org.gradle.java.installations.paths",
        "Keep wrapper, SDK and toolchain versions aligned with repo",
    ], "2:15. Open gradle-daemon-jvm.properties and app/build.gradle.kts. "
       "A generic setup-java 17 example does not describe this project. Downloaded daemon toolchains are not a substitute for the test launcher.")

    code_slide(deck, "Fast job | Build and test on every PR", '''permissions:
  contents: read

steps:
  - uses: actions/checkout@v4
  # Install both JDKs, SDK 37 and Gradle (full script).
  - run: >-
      ./gradlew :app:testDebugUnitTest
      :app:lintDebug :app:assembleDebug
  - uses: actions/upload-artifact@v4
    if: always()
    # Upload reports even when a test fails.''', "2:20. This is an excerpt, not a standalone workflow. "
        "Use complete YAML from script; show push, pull_request, manual trigger and concurrency. No continue-on-error.")

    code_slide(deck, "Device job | Instrumented tests on Linux", '''- name: Run device tests
  uses: reactivecircus/android-emulator-runner@v2
  with:
    api-level: 35
    arch: x86_64
    target: google_apis
    disable-animations: true
    emulator-options: >-
      -no-window -gpu swiftshader_indirect
      -noaudio -no-boot-anim -no-snapshot
    script: ./gradlew :app:connectedDebugAndroidTest''', "2:25. Full workflow includes KVM access, SDK packages, timeouts and report upload. "
        "API 35 satisfies minSdk 26 but is not a target SDK 37 coverage matrix. An extra API is homework.")

    bullet_slide(deck, "Actions lab | Watch the evidence", [
        "Copy the full workflow to .github/workflows/android-tests.yml",
        "Push; inspect JVM and emulator job logs",
        "Locate test XML/HTML, lint report and debug APK artifacts",
        "If red: identify build, toolchain, assertion or device problem",
        "Keep tests failing visibly; do not hide errors with retries",
        "Set required checks after their first completed run",
    ], "2:30-2:45. Run the instructor backup workflow if student infrastructure is slow. "
       "Upload artifacts with always(), but let the failed test step keep the job red.")

    bullet_slide(deck, "Failure drill | Prove the quality gate", [
        "Change quantity increment from +1 to +2 on the feature branch",
        "Run repeated-add test: expected Cart (2), actual Cart (4)",
        "Commit and push; UI tests must turn red",
        "Inspect the failing artifact and blocked merge requirement",
        "Restore +1, rerun, push the fix; checks must return green",
    ], "2:45-2:53. Use the named repeatedAdd test. Do not break main. "
       "If CI was not preflighted, distinguish intended demonstration from observed evidence.")

    bullet_slide(deck, "CI is not deployment", [
        "CI: validate each change with builds and tests",
        "Today's output: unsigned-for-release debug APK artifact",
        "Continuous delivery: approved, signed release available",
        "Continuous deployment: automatically publish approved release",
        "Release bonus: protected environment, approval and secrets",
        "Never put signing keys in Git or expose them to fork PRs",
    ], "2:53. A debug APK download is not Play Store deployment. "
       "Discuss least privilege, action SHA pinning, trusted release branches and restricted signing access.")

    bullet_slide(deck, "Debugging red builds", [
        "Missing SDK/license: inspect setup-android and sdkmanager",
        "No Java 21 launcher: inspect installations paths and JDK logs",
        "Node not found: route, semantics merge, scroll, duplicate tags",
        "State wrong: inspect quantity updates and state lifetime",
        "Emulator crash: KVM, boot logs and runner resource limits",
        "Green locally, red CI: locale, ordering, timing or environment",
    ], "2:56. Most cheap diagnoses come from reading the first failing task. "
       "Our price formatter uses US locale explicitly; test fixtures are recreated per Activity.")

    bullet_slide(deck, "Exit ticket and homework", [
        "Show the Shop -> Cart -> Back journey and its assertions",
        "Explain why UI tests need an emulator job",
        "Show one red regression run and its green fix",
        "Submit repository URL, PR and Actions run evidence",
        "Homework: system Back, recreation, second API level",
        "Bonus: fake async catalog and approved release design",
    ], "2:58-3:00. Apply script rubric: screens 25, UI tests 35, CI 30, explanation 10. "
       "Finish with one learner explaining a failure artifact; do not spend exit ticket time waiting for emulator boot.", True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUTPUT)
    print(f"Saved {OUTPUT} with {len(deck.slides)} slides")


if __name__ == "__main__":
    build_deck()