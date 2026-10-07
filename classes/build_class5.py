"""Generate a concepts-only CI/CD presentation using the established theme."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from build_class4 import bullet_slide, flow_slide


OUTPUT = Path(__file__).resolve().parents[1] / "slides" / "Class5_CI_CD_Concepts.pptx"

CONCEPTS = [
    ("CI/CD", [
        "Continuous integration, delivery and deployment",
        "From a source change to a trusted software artifact",
        "Automation, feedback and controlled release decisions",
    ], "CI/CD combines engineering practices and automation. Its purpose is reliable feedback and repeatable releases, not merely adding a workflow file."),
    ("Why CI/CD exists", [
        "Manual builds are difficult to reproduce consistently",
        "Large, infrequent integrations hide conflicts and regressions",
        "Late feedback makes failures harder to isolate",
        "Repeatable pipelines make results and decisions visible",
    ], "Automation reduces variation, but it does not remove the need for review, good tests or sound release judgment."),
    ("Continuous integration", [
        "Integrate small changes frequently into shared source control",
        "Automatically build and test changes",
        "Give contributors fast, repeatable feedback",
        "Keep the shared branch healthy through enforced checks",
    ], "CI is a working practice supported by a pipeline. Long-lived branches with occasional automated builds are not the full practice of continuous integration."),
    ("Continuous delivery", [
        "Keep passing changes ready for release",
        "Produce a traceable artifact through a repeatable process",
        "A human or business decision authorizes production release",
        "Readiness includes configuration, signing and release policy",
    ], "A test report or debug artifact alone does not establish production release readiness. Continuous delivery retains a deliberate release decision."),
    ("Continuous deployment", [
        "Passing changes proceed automatically to the target environment",
        "No separate manual release decision for each change",
        "Requires strong gates, monitoring and recovery mechanisms",
        "Store review and platform policies can constrain deployment",
    ], "Continuous deployment extends automation through deployment. A package upload is not necessarily activation for users, especially with mobile store review."),
    ("CI, delivery and deployment", [
        "CI answers: does this change integrate and pass our checks?",
        "Delivery answers: is the change ready to release?",
        "Deployment answers: is the change running for its users?",
        "The key delivery/deployment distinction is release authorization",
    ], "Both forms of CD automate preparation. Continuous delivery retains an explicit decision; continuous deployment automates that decision under policy."),
    ("GitHub Actions vocabulary", [
        "Workflow: an event-triggered automation definition",
        "Job: a set of steps on a runner",
        "Step: an action or command inside a job",
        "Runner: the machine executing the job",
        "Action: a reusable automation component",
    ], "Workflow definitions are discovered under .github/workflows. Steps execute in order within a job; jobs are independent unless dependencies are declared."),
    ("Triggers and dependencies", [
        "Push and pull request events validate source changes",
        "Manual dispatch supports explicitly requested runs",
        "Completion events can connect validation and delivery",
        "Job dependencies enforce order and success requirements",
        "Concurrency policies prevent obsolete runs wasting resources",
    ], "Event filters and job conditions define which changes are eligible. Separate runs must use explicit provenance rather than assuming the latest artifact is the correct one."),
    ("Quality gates", [
        "Tests, static analysis and build checks provide evidence",
        "Required checks can prevent merging a failing change",
        "Review adds judgment that automated checks cannot replace",
        "A gate is effective only when configured and enforced",
    ], "Branch protection and rulesets are repository policies, separate from workflow definitions. Their availability and bypass rules affect the actual guarantee."),
    ("Test layers in an Android pipeline", [
        "JUnit: business logic on a host JVM",
        "Robolectric: Android behavior simulated on a host JVM",
        "Instrumented tests: behavior on an emulator or device",
        "Build checks: packaging and compilation, not user behavior",
        "Focused layers balance speed, fidelity and maintenance cost",
    ], "Host tests do not establish device behavior automatically. Existing Shop-independent tests do not cover Shop navigation simply because they run in the same repository."),
    ("Reproducible build environments", [
        "Commit the build wrapper and dependency declarations",
        "Align Java, Gradle, plugins and Android SDK versions",
        "Separate compilation SDK from device runtime API",
        "Make assumptions explicit instead of relying on a laptop",
    ], "In this project the daemon requests Java 25 and host tests request Java 21. SDK 37 compilation and API 35 device testing serve different purposes."),
    ("Artifacts and provenance", [
        "An artifact is a retained output: APK, report or package",
        "Associate outputs with a commit and pipeline run",
        "Promote the tested artifact rather than rebuilding another one",
        "Retention and integrity determine later availability and trust",
    ], "Artifacts and caches have different roles. Artifact promotion preserves the exact tested package; a fresh rebuild can introduce different dependencies or source inputs."),
    ("Feedback and failure diagnosis", [
        "A failed assertion indicates a checked behavior disagrees",
        "Build errors indicate compilation or dependency problems",
        "Runner errors indicate provisioning or infrastructure problems",
        "Logs and reports identify the earliest actionable failure",
        "Preserving evidence must not conceal failed status",
    ], "A red pipeline is a signal requiring classification. Unconditional retries or ignored failures can hide regressions instead of diagnosing them."),
    ("Approvals and environments", [
        "An environment groups delivery policy and scoped access",
        "Reviewers can authorize a protected delivery job",
        "Branch restrictions limit eligible sources",
        "An environment name alone does not enforce approval",
    ], "Protection rules must actually exist. Without reviewers, an environment-bound job may run automatically; repository plan and visibility can limit protection features."),
    ("Pipeline security", [
        "Use the minimum token permissions required",
        "Treat pull-request code and downloaded artifacts as untrusted",
        "Keep credentials and signing keys outside source control",
        "Restrict privileged delivery to trusted source events",
        "Review third-party actions and pin approved commit versions",
    ], "Completion-triggered workflows can hold elevated permissions. They must not execute untrusted downloaded code or expose release secrets to fork pull requests."),
    ("Android distribution boundaries", [
        "Debug APK: installable test build with debug signing",
        "Production APK or AAB: release configuration and signing",
        "Artifact storage is not Play Store publishing",
        "Publishing is not necessarily immediate user rollout",
    ], "The project's internal debug download illustrates tested-artifact delivery, not full production continuous delivery. An Android App Bundle is a publishing format, not directly installed like an APK."),
    ("Caching, speed and reliability", [
        "Caches reduce repeated setup and dependency downloads",
        "Fast host checks and slower device checks can run in parallel",
        "Timeouts bound infrastructure failures",
        "Reliable fixtures improve feedback more than hidden retries",
    ], "Caches are optimizations, not durable delivery artifacts. Parallel checks provide faster feedback, while delivery still waits for all required validation."),
    ("Release risk and recovery", [
        "Staged rollouts limit the impact of an unexpected defect",
        "Monitoring reveals problems beyond pre-release tests",
        "Recovery can require a forward fix with a new version",
        "Mobile signing and version rules constrain downgrades",
    ], "An old APK download is not a complete rollback strategy. Installed application state, signing identity and version codes affect what recovery is possible."),
    ("What a green pipeline means", [
        "The configured checks passed for a particular revision",
        "It does not prove all behavior is correct",
        "Coverage gaps remain explicit engineering risks",
        "Trust comes from tests, provenance, policy and observation",
    ], "A green run is bounded evidence. CI/CD improves confidence through repeatability, but cannot prove the absence of defects."),
]


def build_deck():
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    for index, (title, lines, notes) in enumerate(CONCEPTS):
        bullet_slide(deck, title, lines, notes, dark=index == 0)
        if index == 5:
            flow_slide(deck, "From change to delivery", [
                ("Source", "A traceable revision\nShared version control"),
                ("Validation", "Build and tests\nRequired quality gates"),
                ("Artifact", "Immutable output\nCommit and run identity"),
                ("Delivery", "Release policy\nApproval or automation"),
            ], "The lifecycle connects source, evidence and delivery. The release decision separates continuous delivery from continuous deployment.")
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text == "Class 4 | UI Testing and CI/CD":
                shape.text_frame.paragraphs[0].runs[0].text = "CI/CD | Concepts"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUTPUT)
    print(f"Saved {OUTPUT} with {len(deck.slides)} concepts-only slides")


if __name__ == "__main__":
    build_deck()