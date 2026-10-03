"""Generate the Class 1 slide deck: Introduction to Unit, Integration and UI Tests."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- Theme -------------------------------------------------------------------
NAVY = RGBColor(0x0D, 0x1B, 0x2A)
BLUE = RGBColor(0x1B, 0x4D, 0x7A)
ACCENT = RGBColor(0x3D, 0xDC, 0x84)  # Android green
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
GREY = RGBColor(0x5A, 0x6A, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG = RGBColor(0x1E, 0x1E, 0x2E)
CODE_FG = RGBColor(0xE6, 0xE6, 0xE6)

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def box(slide, l, t, w, h):
    return slide.shapes.add_textbox(l, t, w, h).text_frame


def set_run(run, text, size, color, bold=False, italic=False, font=BODY_FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def accent_bar(slide):
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.18), SH)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def title_bar(slide, number, title):
    """Standard content-slide header."""
    accent_bar(slide)
    tf = box(slide, Inches(0.6), Inches(0.35), Inches(11.8), Inches(1.0))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, title, 32, NAVY, bold=True)
    # slide number chip
    chip = slide.shapes.add_shape(1, Inches(12.3), Inches(0.4), Inches(0.7), Inches(0.55))
    chip.fill.solid()
    chip.fill.fore_color.rgb = BLUE
    chip.line.fill.background()
    ctf = chip.text_frame
    ctf.word_wrap = False
    cp = ctf.paragraphs[0]
    cp.alignment = PP_ALIGN.CENTER
    cr = cp.add_run()
    set_run(cr, str(number), 18, WHITE, bold=True)
    # underline
    ln = slide.shapes.add_shape(1, Inches(0.6), Inches(1.3), Inches(2.4), Inches(0.06))
    ln.fill.solid()
    ln.fill.fore_color.rgb = ACCENT
    ln.line.fill.background()


def bullets(slide, items, top=Inches(1.7), left=Inches(0.8), width=Inches(11.7),
            height=Inches(5.2), size=20):
    """items: list of (text, level) or str."""
    tf = box(slide, left, top, width, height)
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(10)
        bullet = "•  " if level == 0 else "–  "
        r = p.add_run()
        set_run(r, bullet + text, size - level * 2, NAVY if level == 0 else GREY)
    return tf


def speaker(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def code_slide(number, title, code):
    slide = add_slide()
    bg(slide, NAVY)
    accent_bar(slide)
    tf = box(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.9))
    r = tf.paragraphs[0].add_run()
    set_run(r, title, 28, WHITE, bold=True)
    chip = slide.shapes.add_shape(1, Inches(12.3), Inches(0.35), Inches(0.7), Inches(0.55))
    chip.fill.solid(); chip.fill.fore_color.rgb = ACCENT; chip.line.fill.background()
    cr = chip.text_frame.paragraphs[0].add_run()
    chip.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_run(cr, str(number), 18, NAVY, bold=True)
    # code panel
    panel = slide.shapes.add_shape(1, Inches(0.6), Inches(1.35), Inches(12.1), Inches(5.7))
    panel.fill.solid(); panel.fill.fore_color.rgb = CODE_BG
    panel.line.color.rgb = BLUE; panel.line.width = Pt(1)
    ctf = panel.text_frame
    ctf.word_wrap = True
    ctf.vertical_anchor = MSO_ANCHOR.TOP
    ctf.margin_left = Inches(0.3); ctf.margin_top = Inches(0.2)
    ctf.margin_right = Inches(0.2); ctf.margin_bottom = Inches(0.2)
    lines = code.split("\n")
    for i, line in enumerate(lines):
        p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
        p.line_spacing = 1.0
        r = p.add_run()
        set_run(r, line if line else " ", 14, CODE_FG, font=CODE_FONT)
    return slide


def table_slide(number, title, headers, rows, col_widths=None, note=None):
    slide = add_slide()
    bg(slide, WHITE)
    title_bar(slide, number, title)
    nrows = len(rows) + 1
    ncols = len(headers)
    total_w = Inches(11.9)
    left = Inches(0.7)
    top = Inches(1.7)
    height = Inches(0.6) * nrows
    gfx = slide.shapes.add_table(nrows, ncols, left, top, total_w, height)
    table = gfx.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    # header
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); set_run(r, h, 16, WHITE, bold=True)
    # body
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
            p = cell.text_frame.paragraphs[0]
            r = p.add_run(); set_run(r, val, 13, NAVY)
    if note:
        tf = box(slide, Inches(0.7), Inches(6.7), Inches(11.9), Inches(0.6))
        r = tf.paragraphs[0].add_run()
        set_run(r, note, 14, GREY, italic=True)
    return slide


# ---- Diagram helpers ---------------------------------------------------------
def shp(slide, kind, l, t, w, h, fill, line=None, line_w=1.0):
    sp = slide.shapes.add_shape(kind, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is not None:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def label(sp, text, size, color, bold=True, sub=None, sub_color=None):
    tf = sp.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); set_run(r, text, size, color, bold=bold)
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); set_run(r2, sub, size - 6, sub_color or color)


def aaa_flow(slide, top=Inches(4.4)):
    """Arrange -> Act -> Assert horizontal flow."""
    steps = [("Arrange", "set up data", BLUE, WHITE),
             ("Act", "run the code", ACCENT, NAVY),
             ("Assert", "check result", NAVY, WHITE)]
    x = Inches(1.2); w = Inches(3.1); gap = Inches(0.7); h = Inches(1.5)
    for i, (t, sub, c, tc) in enumerate(steps):
        b = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, top, w, h, c)
        label(b, t, 24, tc, sub=sub, sub_color=tc)
        if i < 2:
            shp(slide, MSO_SHAPE.RIGHT_ARROW,
                Emu(int(x) + int(w) + Inches(0.08)), Emu(int(top) + int(Inches(0.5))),
                Emu(int(gap) - int(Inches(0.16))), Inches(0.5), GREY)
        x = Emu(int(x) + int(w) + int(gap))


def unit_icon(slide, left=Inches(8.6), top=Inches(2.3)):
    """A single self-contained component."""
    outer = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.4), Inches(3.2), LIGHT, BLUE, 2)
    inner = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                Emu(int(left) + int(Inches(0.6))), Emu(int(top) + int(Inches(0.9))),
                Inches(2.2), Inches(1.4), BLUE)
    label(inner, "LoginValidator", 16, WHITE)
    cap = box(slide, left, Emu(int(top) + int(Inches(3.25))), Inches(3.4), Inches(0.5))
    cap.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = cap.paragraphs[0].add_run(); set_run(r, "one class, tested alone", 14, GREY, italic=True)


def integration_diagram(slide, left=Inches(7.6), top=Inches(2.1)):
    """Validator + UserStore feeding into LoginService."""
    v = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.3), Inches(1.0), BLUE)
    label(v, "LoginValidator", 14, WHITE)
    u = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, Emu(int(top) + int(Inches(1.5))),
            Inches(2.3), Inches(1.0), BLUE)
    label(u, "UserStore", 14, WHITE)
    svc = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              Emu(int(left) + int(Inches(3.1))), Emu(int(top) + int(Inches(0.75))),
              Inches(2.4), Inches(1.0), ACCENT)
    label(svc, "LoginService", 15, NAVY)
    ax = Emu(int(left) + int(Inches(2.35)))
    aw = Inches(0.65)
    shp(slide, MSO_SHAPE.RIGHT_ARROW, ax, Emu(int(top) + int(Inches(0.3))), aw, Inches(0.4), GREY)
    shp(slide, MSO_SHAPE.RIGHT_ARROW, ax, Emu(int(top) + int(Inches(1.7))), aw, Inches(0.4), GREY)


def phone_mockup(slide, left=Inches(9.2), top=Inches(1.9)):
    """Simple device with a screen, field and button."""
    body = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.7), Inches(4.6), NAVY)
    body.adjustments[0] = 0.08
    screen = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                 Emu(int(left) + int(Inches(0.22))), Emu(int(top) + int(Inches(0.35))),
                 Inches(2.26), Inches(3.9), WHITE)
    fx = Emu(int(left) + int(Inches(0.45)))
    shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, fx, Emu(int(top) + int(Inches(0.85))),
        Inches(1.8), Inches(0.5), LIGHT, GREY, 1)
    shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, fx, Emu(int(top) + int(Inches(1.6))),
        Inches(1.8), Inches(0.5), LIGHT, GREY, 1)
    btn = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, fx, Emu(int(top) + int(Inches(2.45))),
              Inches(1.8), Inches(0.55), ACCENT)
    label(btn, "Log in", 14, NAVY)


def warning_icon(slide, left=Inches(9.4), top=Inches(2.2)):
    tri = shp(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, left, top, Inches(3.0), Inches(2.6),
              RGBColor(0xF4, 0xB4, 0x00))
    tf = tri.text_frame
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(60)
    r = p.add_run(); set_run(r, "!", 54, NAVY, bold=True)
    cap = box(slide, left, Emu(int(top) + int(Inches(2.7))), Inches(3.0), Inches(0.9))
    cp = cap.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
    r = cp.add_run(); set_run(r, "Method ... not mocked", 15, GREY, italic=True, font=CODE_FONT)


def espresso_flow(slide, top=Inches(4.9)):
    steps = [("onView", "matcher", BLUE), ("perform", "action", ACCENT), ("check", "assertion", NAVY)]
    x = Inches(1.3); w = Inches(3.2); h = Inches(1.2)
    for i, (t, sub, c) in enumerate(steps):
        b = shp(slide, MSO_SHAPE.CHEVRON, x, top, w, h, c)
        tc = NAVY if c == ACCENT else WHITE
        label(b, t, 20, tc, sub=sub, sub_color=tc)
        x = Emu(int(x) + int(w) - int(Inches(0.35)))


def semantics_tree(slide, left=Inches(8.1), top=Inches(1.9)):
    root = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
               Emu(int(left) + int(Inches(1.3))), top, Inches(2.2), Inches(0.8), NAVY)
    label(root, "Column", 15, WHITE)
    kids = [("TextField", BLUE), ("Button", ACCENT), ("List", BLUE)]
    kx = left; ky = Emu(int(top) + int(Inches(1.7))); kw = Inches(1.55)
    centres = []
    for t, c in kids:
        b = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, kx, ky, kw, Inches(0.75), c)
        label(b, t, 12, NAVY if c == ACCENT else WHITE)
        centres.append(Emu(int(kx) + int(kw) // 2))
        kx = Emu(int(kx) + int(kw) + int(Inches(0.35)))
    # connectors from root to each child
    rc = Emu(int(left) + int(Inches(1.3)) + int(Inches(1.1)))
    for cxx in centres:
        conn = slide.shapes.add_connector(2, rc, Emu(int(top) + int(Inches(0.8))),
                                          cxx, ky)
        conn.line.color.rgb = GREY; conn.line.width = Pt(1.5)


# ============================================================================
# Slide 1 — Title
# ============================================================================
s = add_slide()
bg(s, NAVY)
# accent block
blk = s.shapes.add_shape(1, 0, Inches(2.7), SW, Inches(0.12))
blk.fill.solid(); blk.fill.fore_color.rgb = ACCENT; blk.line.fill.background()
tf = box(s, Inches(1.0), Inches(1.6), Inches(11.3), Inches(1.4))
r = tf.paragraphs[0].add_run()
set_run(r, "Testing Android Apps", 54, WHITE, bold=True)
tf2 = box(s, Inches(1.0), Inches(2.95), Inches(11.3), Inches(0.9))
r = tf2.paragraphs[0].add_run()
set_run(r, "Class 1 — Introduction to Unit, Integration and UI Tests", 26, ACCENT, bold=True)
tf3 = box(s, Inches(1.0), Inches(4.1), Inches(11.3), Inches(1.2))
for i, line in enumerate(["Tools: JUnit  ·  Espresso  ·  Robolectric  ·  Compose testing"]):
    p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
    rr = p.add_run(); set_run(rr, line, 20, LIGHT)
speaker(s, "Welcome. Today is the foundation. Every later class builds on these three test types.")

# ============================================================================
# Slide 2 — Course roadmap
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 2, "Course roadmap")
bullets(s, [
    ("1.  Unit, integration and UI tests (today)", 0),
    ("2.  TDD", 0),
    ("3.  Mocks, stubs and fakes", 0),
    ("4.  UI testing and CI/CD", 0),
    ("5.  Fixing flaky tests", 0),
    ("6.  Catching bugs, plus the exam", 0),
])
speaker(s, "Today is the foundation. Every later class builds on these three test types.")

# ============================================================================
# Slide 3 — Why write tests?
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 3, "Why write tests?")
bullets(s, [
    "Catch regressions before users do",
    "Change code without fear",
    "Tests document how the code should behave",
    "Bugs cost more the later you find them",
])
speaker(s, 'Ask the class: "Who has broken one feature while fixing another?"')

# ============================================================================
# Slide 4 — What a test looks like
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 4, "What a test looks like")
bullets(s, [
    "Arrange: set up the data",
    "Act: run the code",
    "Assert: check the result",
    "Good tests are fast, independent, repeatable and self-checking",
    "Name tests as whatIsTested_condition_expectedResult",
], height=Inches(2.4))
aaa_flow(s, top=Inches(4.6))

# ============================================================================
# Slide 5 — Unit tests
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 5, "Unit tests")
bullets(s, [
    "Test one class or function on its own",
    "Take milliseconds and need no device",
    'Example: "Does LoginValidator reject an empty email?"',
], width=Inches(7.4))
unit_icon(s)

# ============================================================================
# Slide 6 — Integration tests
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 6, "Integration tests")
bullets(s, [
    "Test several real components working together",
    "Catch problems where pieces connect: wrong data passed, wrong order, wrong assumptions",
    'Example: "Does LoginService use the validator and the user store correctly?"',
], width=Inches(6.6), height=Inches(2.0))
integration_diagram(s, top=Inches(4.2), left=Inches(3.7))
speaker(s, "Today we use only real classes. Replacing parts with fakes and mocks is class 3.")

# ============================================================================
# Slide 7 — UI tests
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 7, "UI tests")
bullets(s, [
    "Test from the user's point of view: type, click, see",
    "The most realistic test type, and the slowest and most fragile",
    'Example: "After I tap Log in with no email, do I see \'Email is required\'?"',
], width=Inches(7.6))
phone_mockup(s)

# ============================================================================
# Slide 8 — The testing pyramid
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 8, "The testing pyramid")
bullets(s, [
    "Many unit tests, some integration tests, a few UI tests",
    "A rough guide is 70 / 20 / 10",
    "Higher up means slower, more realistic and more expensive to maintain",
], width=Inches(6.2))
# draw pyramid
cx = Inches(9.7)
levels = [("UI  10%", Inches(2.0), ACCENT), ("Integration  20%", Inches(3.4), BLUE), ("Unit  70%", Inches(4.8), NAVY)]
top = Inches(2.0)
for lvl_text, w, color in levels:
    tri = s.shapes.add_shape(1, Emu(int(cx) - int(w) // 2), top, w, Inches(1.25))
    tri.fill.solid(); tri.fill.fore_color.rgb = color; tri.line.color.rgb = WHITE; tri.line.width = Pt(2)
    tp = tri.text_frame.paragraphs[0]; tp.alignment = PP_ALIGN.CENTER
    rr = tp.add_run(); set_run(rr, lvl_text, 16, WHITE, bold=True)
    top = Emu(int(top) + int(Inches(1.3)))

# ============================================================================
# Slide 9 — Two source sets
# ============================================================================
table_slide(9, "Two source sets in Android",
            ["", "src/test (local)", "src/androidTest (instrumented)"],
            [
                ["Runs on", "Your computer's JVM", "Emulator or device"],
                ["Speed", "Fast", "Slow"],
                ["Android APIs", "Stubbed, or real through Robolectric", "Real"],
                ["Command", "./gradlew testDebugUnitTest", "./gradlew connectedDebugAndroidTest"],
            ],
            col_widths=[Inches(2.4), Inches(4.6), Inches(4.9)])

# ============================================================================
# Slide 10 — JUnit 4 essentials
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 10, "JUnit 4 essentials")
bullets(s, [
    "@Test, @Before, @After, @get:Rule",
    "assertEquals, assertTrue, assertFalse, assertNull, assertThrows",
    "@RunWith changes which runner executes the test",
])

# ============================================================================
# Slide 11 — JUnit example (code)
# ============================================================================
code_slide(11, "JUnit example", '''class CalculatorTest {
    private lateinit var calc: Calculator

    @Before fun setUp() { calc = Calculator() }            // Arrange

    @Test fun add_twoPositives_returnsSum() {
        val result = calc.add(2, 3)                         // Act
        assertEquals(5, result)                             // Assert
    }

    @Test fun divide_byZero_throws() {
        assertThrows(ArithmeticException::class.java) {
            calc.divide(1, 0)
        }
    }
}''')

# ============================================================================
# Slide 12 — The Android problem
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 12, "The Android problem")
bullets(s, [
    "Local tests compile against a stubbed android.jar",
    'Android calls fail with "Method ... not mocked", or return null',
    "Two fixes: keep logic free of Android code (best), or use Robolectric",
], width=Inches(7.6))
warning_icon(s)

# ============================================================================
# Slide 13 — Robolectric (bullets + code)
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 13, "Robolectric")
bullets(s, [
    "Runs real Android framework code on the JVM",
    "Can start Activities, inflate layouts and read resources",
    "Runs in seconds with no emulator",
    "Its limit: it isn't a real device (no real rendering or hardware)",
], height=Inches(3.2))
panel = s.shapes.add_shape(1, Inches(0.8), Inches(5.1), Inches(9.5), Inches(1.7))
panel.fill.solid(); panel.fill.fore_color.rgb = CODE_BG; panel.line.color.rgb = BLUE
ctf = panel.text_frame; ctf.margin_left = Inches(0.25); ctf.margin_top = Inches(0.15)
for i, line in enumerate(["@RunWith(AndroidJUnit4::class)", "@Config(sdk = [35])", "class MyTest { ... }"]):
    p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
    r = p.add_run(); set_run(r, line, 15, CODE_FG, font=CODE_FONT)

# ============================================================================
# Slide 14 — Espresso
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 14, "Espresso")
bullets(s, [
    "The UI testing framework for XML (View-based) screens",
    "Formula: onView(matcher).perform(action).check(assertion)",
    "Waits for the UI thread to be idle, so you never need Thread.sleep()",
    "Launch a screen with ActivityScenarioRule(MyActivity::class.java)",
], height=Inches(2.7))
espresso_flow(s, top=Inches(5.0))

# ============================================================================
# Slide 15 — Espresso cheat sheet
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 15, "Espresso cheat sheet")
bullets(s, [
    ("Matchers", 0),
    ("withId, withText, withHint, isDisplayed, isEnabled", 1),
    ("Actions", 0),
    ("click(), typeText(), replaceText(), closeSoftKeyboard(), scrollTo()", 1),
    ("Assertions", 0),
    ("matches(...), doesNotExist()", 1),
])

# ============================================================================
# Slide 16 — Testing Compose
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 16, "Testing Compose")
bullets(s, [
    "Composables aren't Views, so Espresso's onView can't find them",
    "Compose tests read the semantics tree — meaning rather than pixels",
    "createComposeRule() tests a single composable",
    "createAndroidComposeRule<MainActivity>() tests a whole Activity",
    "The rule uses Espresso internally to wait for the UI",
], width=Inches(7.0), height=Inches(4.0))
semantics_tree(s, left=Inches(8.1), top=Inches(3.1))

# ============================================================================
# Slide 17 — Compose testing cheat sheet
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 17, "Compose testing cheat sheet")
bullets(s, [
    ("Finders", 0),
    ("onNodeWithText, onNodeWithTag, onNodeWithContentDescription", 1),
    ("Actions", 0),
    ("performClick(), performTextInput(), performScrollTo()", 1),
    ("Assertions", 0),
    ("assertIsDisplayed(), assertIsEnabled(), assertIsNotEnabled(), assertDoesNotExist()", 1),
    ('Label nodes with Modifier.testTag("id"), similar to R.id in XML', 0),
], size=19)

# ============================================================================
# Slide 18 — Choosing a tool (table)
# ============================================================================
table_slide(18, "Choosing a tool",
            ["What you test", "Test type", "Tool", "Source set"],
            [
                ["One Kotlin class", "Unit", "JUnit", "test"],
                ["Several classes together", "Integration", "JUnit (+ Robolectric if Android)", "test"],
                ["XML screen, fast", "UI", "Espresso on Robolectric", "test"],
                ["XML screen, realistic", "UI", "Espresso", "androidTest"],
                ["Compose screen, fast", "UI", "Compose test on Robolectric", "test"],
                ["Compose screen, realistic", "UI", "Compose test", "androidTest"],
            ],
            col_widths=[Inches(3.2), Inches(2.0), Inches(4.3), Inches(2.4)])

# ============================================================================
# Slide 19 — Running tests
# ============================================================================
s = add_slide(); bg(s, WHITE); title_bar(s, 19, "Running tests")
bullets(s, [
    "In Android Studio, click the green ▶ next to a test or class",
    "./gradlew testDebugUnitTest",
    "./gradlew connectedDebugAndroidTest",
    "Reports are in app/build/reports/tests/ and app/build/reports/androidTests/",
])

# ============================================================================
# Slide 20 — Summary and exercises
# ============================================================================
s = add_slide(); bg(s, NAVY); accent_bar(s)
tf = box(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.0))
r = tf.paragraphs[0].add_run(); set_run(r, "Summary and exercises", 32, WHITE, bold=True)
ln = s.shapes.add_shape(1, Inches(0.65), Inches(1.35), Inches(2.4), Inches(0.06))
ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()
tf = box(s, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0)); tf.word_wrap = True
items = [
    "Unit tests check one piece, integration tests check pieces together, UI tests check what the user sees",
    "JUnit is the base, Robolectric brings Android to the JVM, Espresso and Compose tests drive the UI",
    "Exercise 1: XML Login screen",
    "Exercise 2: Compose Todo screen",
]
for i, it in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(14)
    color = ACCENT if it.startswith("Exercise") else LIGHT
    r = p.add_run(); set_run(r, "•  " + it, 22, color, bold=it.startswith("Exercise"))

prs.save("slides/Class1_Testing_Android.pptx")
print("Saved slides/Class1_Testing_Android.pptx with", len(prs.slides._sldIdLst), "slides")
